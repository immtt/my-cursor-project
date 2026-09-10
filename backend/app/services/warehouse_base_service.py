"""仓库主数据 `warehouse_base`：与「仓管理」Excel 列对齐，另含 `group_name`（集团）。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook

from sqlalchemy import and_, or_, select
from sqlalchemy.orm import Session, selectinload

from app.models.entities import WarehouseBase, WarehouseBusinessBrand
from app.utils.finite import json_safe_float, parse_finite_or_none

# 导出表头顺序：在「仓管理」模板基础上，于「仓库名称」后插入「集团」
_EXCEL_HEADERS: List[Tuple[str, str]] = [
    ("仓库代码", "warehouse_code"),
    ("仓库名称", "warehouse_name"),
    ("集团", "group_name"),
    ("所属组织", "owning_org"),
    ("业务品牌", "brand"),
    ("物流组织", "logistics_org"),
    ("配送中心门店名称", "dc_store_name"),
    ("仓库分类", "warehouse_category"),
    ("仓库温层", "temperature_layer"),
    ("仓库负责人", "manager_name"),
    ("联系电话", "manager_phone"),
    ("仓库地址", "address"),
    ("仓库坐标", "coordinate_raw"),
    ("状态", "status"),
    ("仓库类型", "warehouse_type"),
    ("仓库经营类型", "business_type"),
    ("仓库产权", "property_type"),
    ("收货联系人", "receiver_contact"),
    ("收货电话", "receiver_phone"),
    ("仓库面积", "area_sqm"),
    ("覆盖门店区域", "coverage_region"),
    ("库区功能", "zone_function"),
    ("预计覆盖门店数", "expected_store_count"),
    ("本月真实覆盖门店数量", "monthly_covered_stores"),
    ("开仓日", "opening_date"),
    ("覆盖SKU数", "sku_count_text"),
    ("是否启用采购共享仓", "purchase_shared_flag"),
    ("是否启用采购直通", "purchase_direct_flag"),
    ("当前仓是否商品组下单仓", "is_group_order_warehouse"),
    ("备注", "remark"),
    ("申请人", "applicant"),
    ("创建时间", "source_created_at"),
]

def _dt_iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def _cell_str(v: Any) -> Optional[str]:
    if v is None:
        return None
    if isinstance(v, float):
        finite = parse_finite_or_none(v)
        if finite is None:
            return None
        if finite == int(finite):
            return str(int(finite))
        v = finite
    s = str(v).strip()
    return s if s else None


def _cell_float(v: Any) -> Optional[float]:
    return parse_finite_or_none(v)


def business_brand_to_item(b: WarehouseBusinessBrand) -> Dict[str, Any]:
    return {
        "id": b.id,
        "name": b.name,
        "logo_as": b.logo_as,
        "logo_url": f"/api/warehouse-base/business-brands/{b.id}/file" if b.logo_path else None,
    }


def warehouse_base_to_item(row: WarehouseBase) -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "id": row.id,
        "warehouse_code": row.warehouse_code,
        "warehouse_name": row.warehouse_name,
        "group_name": row.group_name,
        "owning_org": row.owning_org,
        "brand": row.brand,
        "logistics_org": row.logistics_org,
        "dc_store_name": row.dc_store_name,
        "warehouse_category": row.warehouse_category,
        "temperature_layer": row.temperature_layer,
        "manager_name": row.manager_name,
        "manager_phone": row.manager_phone,
        "address": row.address,
        "coordinate_raw": row.coordinate_raw,
        "status": row.status,
        "warehouse_type": row.warehouse_type,
        "business_type": row.business_type,
        "property_type": row.property_type,
        "receiver_contact": row.receiver_contact,
        "receiver_phone": row.receiver_phone,
        "area_sqm": json_safe_float(row.area_sqm),
        "coverage_region": row.coverage_region,
        "zone_function": row.zone_function,
        "expected_store_count": json_safe_float(row.expected_store_count),
        "monthly_covered_stores": row.monthly_covered_stores,
        "opening_date": row.opening_date,
        "sku_count_text": row.sku_count_text,
        "purchase_shared_flag": row.purchase_shared_flag,
        "purchase_direct_flag": row.purchase_direct_flag,
        "is_group_order_warehouse": row.is_group_order_warehouse,
        "remark": row.remark,
        "applicant": row.applicant,
        "source_created_at": row.source_created_at,
        "import_source": row.import_source,
        "created_at": _dt_iso(row.created_at),
        "updated_at": _dt_iso(row.updated_at),
    }
    ch = getattr(row, "business_brands", None) or []
    out["business_brands"] = [business_brand_to_item(b) for b in sorted(ch, key=lambda x: (x.sort_order, x.id or 0))]
    return out


_STR_ATTRS = (
    "warehouse_name",
    "group_name",
    "owning_org",
    "brand",
    "logistics_org",
    "dc_store_name",
    "warehouse_category",
    "temperature_layer",
    "manager_name",
    "manager_phone",
    "address",
    "coordinate_raw",
    "status",
    "warehouse_type",
    "business_type",
    "property_type",
    "receiver_contact",
    "receiver_phone",
    "coverage_region",
    "zone_function",
    "monthly_covered_stores",
    "opening_date",
    "sku_count_text",
    "purchase_shared_flag",
    "purchase_direct_flag",
    "is_group_order_warehouse",
    "remark",
    "applicant",
    "source_created_at",
    "import_source",
)
_FLOAT_ATTRS = ("area_sqm", "expected_store_count")


def _set_str_attr(row: WarehouseBase, k: str, v: Any) -> None:
    if v is None or v == "":
        setattr(row, k, None)
    else:
        setattr(row, k, str(v).strip() if not isinstance(v, str) else v.strip() or None)


def _set_float_attr(row: WarehouseBase, k: str, v: Any) -> None:
    if v is None or v == "":
        setattr(row, k, None)
        return
    parsed = parse_finite_or_none(v)
    if parsed is not None:
        setattr(row, k, parsed)
        return
    try:
        float(v)
    except (TypeError, ValueError):
        setattr(row, k, None)
        return
    raise ValueError(f"{k} 须为有限数值")


def _apply_geocode_to_warehouse_row(db: Session, row: WarehouseBase) -> None:
    """根据 `address` 调用高德（或模拟）写入 `coordinate_raw`；解析失败则保持原值。"""
    from app.services.gaode_service import geocode_address_to_coordinate_raw

    s = geocode_address_to_coordinate_raw(row.address)
    if s is None:
        return
    row.coordinate_raw = s
    db.add(row)
    db.commit()
    db.refresh(row)


def batch_supplement_warehouse_base_coordinates(
    db: Session,
    *,
    only_missing: bool = True,
    warehouse: str = "",
    group_name: str = "",
    owning_org: str = "",
    brand: str = "",
    status: str = "",
    warehouse_type: str = "",
) -> Dict[str, Any]:
    """
    按与列表接口相同的条件筛选，对命中的仓逐条用「仓库地址」补全/刷新坐标。
    仅当 only_missing 为真时，跳过已填「仓库坐标」的行。
    """
    from app.services.gaode_service import geocode_address_to_coordinate_raw

    q = _warehouse_base_filtered_query(
        db,
        warehouse=warehouse,
        group_name=group_name,
        owning_org=owning_org,
        brand=brand,
        status=status,
        warehouse_type=warehouse_type,
    )
    q = q.order_by(WarehouseBase.id.asc())
    total_matched = q.count()
    cap = 3000
    rows = q.limit(cap).all()
    updated = 0
    failed = 0
    skipped = 0
    for wrow in rows:
        addr = (wrow.address or "").strip()
        if not addr:
            skipped += 1
            continue
        if only_missing and (wrow.coordinate_raw or "").strip():
            skipped += 1
            continue
        s = geocode_address_to_coordinate_raw(addr)
        if s is None:
            failed += 1
            continue
        wrow.coordinate_raw = s
        db.add(wrow)
        updated += 1
    if updated:
        db.commit()
    return {
        "ok": True,
        "only_missing": only_missing,
        "total_matched": total_matched,
        "processed": len(rows),
        "capped": total_matched > cap,
        "updated": updated,
        "failed": failed,
        "skipped": skipped,
    }


def _validate_warehouse_base_core(row: WarehouseBase) -> None:
    if not (row.warehouse_name or "").strip():
        raise ValueError("仓库名称不能为空")
    if not (row.group_name or "").strip():
        raise ValueError("集团不能为空")
    addr = row.address
    if not (addr is not None and str(addr).strip()):
        raise ValueError("仓库地址不能为空")
    if not (row.brand or "").strip():
        raise ValueError("业务品牌不能为空（汇总或至少一条明细）")


def _sync_warehouse_brand_summary(db: Session, row: WarehouseBase) -> None:
    ch = (
        db.query(WarehouseBusinessBrand)
        .filter(WarehouseBusinessBrand.warehouse_id == row.id)
        .order_by(WarehouseBusinessBrand.sort_order, WarehouseBusinessBrand.id)
        .all()
    )
    if not ch:
        return
    parts: List[str] = []
    for x in ch:
        nm = (x.name or "").strip()
        if not nm:
            continue
        la = (x.logo_as or "").strip()
        if la:
            parts.append(f"{nm}（LOGO：{la}）")
        else:
            parts.append(nm)
    row.brand = ";".join(parts) if parts else None


def _replace_business_brands_from_payload(
    db: Session, row: WarehouseBase, data: Dict[str, Any], *, is_create: bool
) -> None:
    items = data.get("business_brands")
    if items is not None:
        if not isinstance(items, list):
            raise ValueError("business_brands 须为数组")
        db.query(WarehouseBusinessBrand).filter(
            WarehouseBusinessBrand.warehouse_id == row.id
        ).delete(synchronize_session=False)
        for i, it in enumerate(items):
            if not isinstance(it, dict):
                continue
            n = str(it.get("name") or "").strip()
            if not n:
                continue
            la = str(it.get("logo_as") or "").strip() or None
            db.add(
                WarehouseBusinessBrand(
                    warehouse_id=row.id,
                    name=n,
                    logo_as=la,
                    sort_order=i,
                )
            )
        return
    if not is_create:
        return
    summary = (row.brand or "").strip()
    if not summary:
        return
    for i, part in enumerate([s.strip() for s in summary.split(";") if s.strip()]):
        db.add(
            WarehouseBusinessBrand(
                warehouse_id=row.id,
                name=part,
                logo_as=None,
                sort_order=i,
            )
        )


def _apply_payload(row: WarehouseBase, data: Dict[str, Any], *, for_create: bool) -> None:
    if for_create:
        code = (data.get("warehouse_code") or "").strip() or None
        name = (data.get("warehouse_name") or "").strip()
        if not name:
            raise ValueError("仓库名称不能为空")
        row.warehouse_code = code
        row.warehouse_name = name
    else:
        if "warehouse_code" in data:
            raw = data["warehouse_code"]
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                row.warehouse_code = None
            else:
                row.warehouse_code = str(raw).strip() or None
        if "warehouse_name" in data and data["warehouse_name"] is not None:
            n = str(data["warehouse_name"]).strip()
            if n:
                row.warehouse_name = n
    for k in _STR_ATTRS:
        if for_create:
            _set_str_attr(row, k, data.get(k))
        elif k in data:
            _set_str_attr(row, k, data[k])
    for k in _FLOAT_ATTRS:
        if for_create:
            _set_float_attr(row, k, data.get(k))
        elif k in data:
            _set_float_attr(row, k, data[k])


def _warehouse_base_filtered_query(
    db: Session,
    *,
    warehouse: str = "",
    group_name: str = "",
    owning_org: str = "",
    brand: str = "",
    status: str = "",
    warehouse_type: str = "",
):
    q = db.query(WarehouseBase)
    wq = (warehouse or "").strip()
    if wq:
        like = f"%{wq}%"
        q = q.filter(
            or_(
                and_(
                    WarehouseBase.warehouse_code.isnot(None),
                    WarehouseBase.warehouse_code.like(like),
                ),
                WarehouseBase.warehouse_name.like(like),
            )
        )

    def _eq(col, raw: str) -> None:
        nonlocal q
        s = (raw or "").strip()
        if s:
            q = q.filter(col == s)

    _eq(WarehouseBase.group_name, group_name)
    _eq(WarehouseBase.owning_org, owning_org)
    bq = (brand or "").strip()
    if bq:
        sub_ids = select(WarehouseBusinessBrand.warehouse_id).where(
            WarehouseBusinessBrand.name == bq
        )
        q = q.filter(
            or_(
                WarehouseBase.brand == bq,
                WarehouseBase.id.in_(sub_ids),
            )
        )
    _eq(WarehouseBase.status, status)
    _eq(WarehouseBase.warehouse_type, warehouse_type)
    return q


def list_warehouse_base_filter_options(db: Session) -> Dict[str, Any]:
    """枚举筛选项 + `warehouses` 列表供「代码+名称」组合框在本地再模糊。"""
    out: Dict[str, Any] = {}
    for col, name in [
        (WarehouseBase.group_name, "group_name"),
        (WarehouseBase.owning_org, "owning_org"),
        (WarehouseBase.brand, "brand"),
        (WarehouseBase.status, "status"),
        (WarehouseBase.warehouse_type, "warehouse_type"),
    ]:
        rows = (
            db.query(col)
            .filter(and_(col.isnot(None), col != ""))
            .distinct()
            .all()
        )
        vals: List[str] = []
        for (raw,) in rows:
            if raw is None:
                continue
            t = str(raw).strip()
            if t:
                vals.append(t)
        out[name] = sorted(set(vals), key=str)

    bset = set(out.get("brand") or [])
    for (bn,) in (
        db.query(WarehouseBusinessBrand.name).filter(WarehouseBusinessBrand.name.isnot(None)).distinct().all()
    ):
        t = (str(bn) if bn is not None else "").strip()
        if t:
            bset.add(t)
    out["brand"] = sorted(bset, key=str)

    pairs: List[Dict[str, Any]] = []
    for c, n in (
        db.query(WarehouseBase.warehouse_code, WarehouseBase.warehouse_name)
        .filter(
            and_(
                WarehouseBase.warehouse_name.isnot(None),
                WarehouseBase.warehouse_name != "",
            )
        )
        .all()
    ):
        ns = (str(n) if n is not None else "").strip()
        if not ns:
            continue
        cs = (str(c) if c is not None else "").strip() or None
        pairs.append({"warehouse_code": cs, "warehouse_name": ns})
    pairs.sort(
        key=lambda x: (x.get("warehouse_code") or "", x.get("warehouse_name") or "")
    )
    out["warehouses"] = pairs
    return out


def list_warehouse_bases(
    db: Session,
    skip: int = 0,
    limit: int = 20,
    *,
    warehouse: str = "",
    group_name: str = "",
    owning_org: str = "",
    brand: str = "",
    status: str = "",
    warehouse_type: str = "",
) -> Tuple[int, List[WarehouseBase]]:
    q = _warehouse_base_filtered_query(
        db,
        warehouse=warehouse,
        group_name=group_name,
        owning_org=owning_org,
        brand=brand,
        status=status,
        warehouse_type=warehouse_type,
    )
    total = q.count()
    rows = (
        q.options(selectinload(WarehouseBase.business_brands))
        .order_by(WarehouseBase.id.desc())
        .offset(max(0, skip))
        .limit(min(200, max(1, limit)))
        .all()
    )
    return total, rows


def get_warehouse_base(db: Session, row_id: int) -> Optional[WarehouseBase]:
    return (
        db.query(WarehouseBase)
        .options(selectinload(WarehouseBase.business_brands))
        .filter(WarehouseBase.id == row_id)
        .first()
    )


def get_by_warehouse_code(db: Session, code: str) -> Optional[WarehouseBase]:
    c = (code or "").strip()
    if not c:
        return None
    return db.query(WarehouseBase).filter(WarehouseBase.warehouse_code == c).first()


def _import_row_dedup_key(
    code: Optional[str], name: str, group_val: str, address_val: str
) -> Tuple[str, str, str, str, str]:
    """同次导入去重用：有仓库代码以代码为键，否则 (名称+集团+地址) 为键。"""
    c = (code or "").strip() or None
    if c:
        return ("c", c, "", "", "")
    return ("g", name, group_val, address_val, "")


def _delete_all_warehouse_bases_for_import_replace(db: Session) -> None:
    """全量覆盖导入前：清空子表后清空主表，使库中仅保留本次成功写入行。"""
    db.query(WarehouseBusinessBrand).delete(synchronize_session=False)
    db.query(WarehouseBase).delete(synchronize_session=False)
    db.flush()


def create_warehouse_base(db: Session, data: Dict[str, Any]) -> WarehouseBase:
    code = (data.get("warehouse_code") or "").strip() or None
    if code and get_by_warehouse_code(db, code):
        raise ValueError("仓库代码已存在")
    row = WarehouseBase(
        warehouse_code=code,
        warehouse_name=(data.get("warehouse_name") or "").strip(),
    )
    _apply_payload(row, data, for_create=True)
    db.add(row)
    db.flush()
    _replace_business_brands_from_payload(db, row, data, is_create=True)
    _sync_warehouse_brand_summary(db, row)
    _validate_warehouse_base_core(row)
    db.commit()
    db.refresh(row)
    _apply_geocode_to_warehouse_row(db, row)
    return get_warehouse_base(db, row.id) or row


def update_warehouse_base(db: Session, row_id: int, data: Dict[str, Any]) -> Optional[WarehouseBase]:
    row = get_warehouse_base(db, row_id)
    if not row:
        return None
    if "warehouse_code" in data:
        raw = data["warehouse_code"]
        c = None if raw is None or (isinstance(raw, str) and not raw.strip()) else str(raw).strip() or None
        if c and c != (row.warehouse_code or None) and get_by_warehouse_code(db, c):
            raise ValueError("仓库代码已存在")
    _apply_payload(row, data, for_create=False)
    if "business_brands" in data:
        _replace_business_brands_from_payload(db, row, data, is_create=False)
    elif "brand" in data:
        db.query(WarehouseBusinessBrand).filter(
            WarehouseBusinessBrand.warehouse_id == row.id
        ).delete(synchronize_session=False)
        summary = (row.brand or "").strip()
        if summary:
            for i, part in enumerate([s.strip() for s in summary.split(";") if s.strip()]):
                db.add(
                    WarehouseBusinessBrand(
                        warehouse_id=row.id,
                        name=part,
                        logo_as=None,
                        sort_order=i,
                    )
                )
    _sync_warehouse_brand_summary(db, row)
    _validate_warehouse_base_core(row)
    db.commit()
    db.refresh(row)
    _apply_geocode_to_warehouse_row(db, row)
    return get_warehouse_base(db, row_id)


def delete_warehouse_base(db: Session, row_id: int) -> bool:
    row = get_warehouse_base(db, row_id)
    if not row:
        return False
    db.delete(row)
    db.commit()
    return True


def _ingest_business_brands_from_import(
    db: Session, w: WarehouseBase, brand_cell: str, logo_as_cell: Optional[str]
) -> None:
    name_parts = [s.strip() for s in (brand_cell or "").split(";") if s.strip()]
    if not name_parts:
        return
    la_parts: List[Optional[str]] = []
    if logo_as_cell and str(logo_as_cell).strip():
        la_parts = [s.strip() or None for s in str(logo_as_cell).split(";")]
    db.query(WarehouseBusinessBrand).filter(WarehouseBusinessBrand.warehouse_id == w.id).delete(
        synchronize_session=False
    )
    for i, nm in enumerate(name_parts):
        la = la_parts[i] if i < len(la_parts) else None
        db.add(
            WarehouseBusinessBrand(warehouse_id=w.id, name=nm, logo_as=la, sort_order=i)
        )
    _sync_warehouse_brand_summary(db, w)


def build_warehouse_base_export_xlsx_bytes(
    db: Session,
    *,
    warehouse: str = "",
    group_name: str = "",
    owning_org: str = "",
    brand: str = "",
    status: str = "",
    warehouse_type: str = "",
) -> bytes:
    q = _warehouse_base_filtered_query(
        db,
        warehouse=warehouse,
        group_name=group_name,
        owning_org=owning_org,
        brand=brand,
        status=status,
        warehouse_type=warehouse_type,
    )
    rows = q.options(selectinload(WarehouseBase.business_brands)).order_by(WarehouseBase.id.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "仓管理"
    headers = [h for h, _ in _EXCEL_HEADERS]
    ws.append(headers)
    attrs = [a for _, a in _EXCEL_HEADERS]
    for r in rows:
        line = []
        for a in attrs:
            v = getattr(r, a)
            if a in ("area_sqm", "expected_store_count"):
                line.append(v if v is not None else "")
            else:
                line.append(v if v is not None else "")
        ws.append(line)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def import_warehouse_workbook(db: Session, path: str, import_source: str) -> Dict[str, Any]:
    # 不用 read_only：部分导出表首行在只读模式下列范围异常，导致表头只读到第一列
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        if "仓管理" in wb.sheetnames:
            ws = wb["仓管理"]
        else:
            ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise ValueError("Excel 无数据行")
    header_row = [str(c).strip() if c is not None else "" for c in rows[0]]
    hmap = {h: i for i, h in enumerate(header_row) if h}
    if "仓库代码" not in hmap or "仓库名称" not in hmap:
        raise ValueError("表头须含「仓库代码」「仓库名称」")

    def pick(row: tuple, cn: str) -> Any:
        idx = hmap.get(cn)
        if idx is None:
            return None
        return row[idx] if idx < len(row) else None

    n_skip = 0
    parsed: List[Tuple[tuple, Optional[str], str, str, str, str]] = []
    for row in rows[1:]:
        code = _cell_str(pick(row, "仓库代码"))
        name = _cell_str(pick(row, "仓库名称"))
        group_val = _cell_str(pick(row, "集团"))
        ad0 = pick(row, "仓库地址")
        address_val = _cell_str(ad0) if ad0 is not None else None
        brand_val = _cell_str(pick(row, "品牌")) or _cell_str(pick(row, "二级组织"))
        if not name or not group_val or not address_val or not brand_val:
            n_skip += 1
            continue
        ckey = (code or "").strip() or None
        parsed.append((row, ckey, name, group_val, address_val, brand_val))

    def _k(parts: Tuple[tuple, Optional[str], str, str, str, str]) -> Tuple[str, str, str, str, str]:
        _r, ckey, name, group_val, address_val, _b = parts
        return _import_row_dedup_key(ckey, name, group_val, address_val)

    deduped: List[Tuple[tuple, Optional[str], str, str, str, str]] = []
    for pr in parsed:
        key = _k(pr)
        deduped = [q for q in deduped if _k(q) != key]
        deduped.append(pr)

    _delete_all_warehouse_bases_for_import_replace(db)
    n_upsert = 0
    for row, ckey, name, group_val, address_val, brand_val in deduped:
        w = WarehouseBase(warehouse_code=ckey, warehouse_name=name)
        w.warehouse_name = name
        w.group_name = group_val
        w.owning_org = _cell_str(pick(row, "所属组织"))
        w.logistics_org = _cell_str(pick(row, "物流组织"))
        w.dc_store_name = _cell_str(pick(row, "配送中心门店名称"))
        w.warehouse_category = _cell_str(pick(row, "仓库分类"))
        w.temperature_layer = _cell_str(pick(row, "仓库温层"))
        w.manager_name = _cell_str(pick(row, "仓库负责人"))
        w.manager_phone = _cell_str(pick(row, "联系电话"))
        w.address = address_val
        if w.address and len(w.address) > 20000:
            w.address = w.address[:20000]
        w.coordinate_raw = _cell_str(pick(row, "仓库坐标"))
        w.status = _cell_str(pick(row, "状态"))
        w.warehouse_type = _cell_str(pick(row, "仓库类型"))
        w.business_type = _cell_str(pick(row, "仓库经营类型"))
        w.property_type = _cell_str(pick(row, "仓库产权"))
        w.receiver_contact = _cell_str(pick(row, "收货联系人"))
        w.receiver_phone = _cell_str(pick(row, "收货电话"))
        w.area_sqm = _cell_float(pick(row, "仓库面积"))
        w.coverage_region = _cell_str(pick(row, "覆盖门店区域"))
        w.zone_function = _cell_str(pick(row, "库区功能"))
        w.expected_store_count = _cell_float(pick(row, "预计覆盖门店数"))
        v_m = pick(row, "本月真实覆盖门店数量")
        w.monthly_covered_stores = _cell_str(v_m) if v_m is not None else None
        w.opening_date = _cell_str(pick(row, "开仓日"))
        w.sku_count_text = _cell_str(pick(row, "覆盖SKU数"))
        w.purchase_shared_flag = _cell_str(pick(row, "是否启用采购共享仓"))
        w.purchase_direct_flag = _cell_str(pick(row, "是否启用采购直通"))
        w.is_group_order_warehouse = _cell_str(pick(row, "当前仓是否商品组下单仓"))
        w.remark = _cell_str(pick(row, "备注"))
        w.applicant = _cell_str(pick(row, "申请人"))
        w.source_created_at = _cell_str(pick(row, "创建时间"))
        w.import_source = (import_source or "").strip() or None
        db.add(w)
        db.flush()
        la_cell = _cell_str(pick(row, "LOGO沿用")) or _cell_str(
            pick(row, "品牌LOGO说明")
        )
        _ingest_business_brands_from_import(db, w, brand_val, la_cell)
        n_upsert += 1
    db.commit()
    return {
        "ok": True,
        "upserted_rows": n_upsert,
        "skipped_incomplete": n_skip,
        "import_source": import_source,
        "replaced_table": True,
    }
