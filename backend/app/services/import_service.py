from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.entities import ManualRoute, SysSuggest


REQUIRED_COLUMNS = [
    "排线日期",
    "运单号",
    "归属线路",
    "始发仓库",
    "拼载门店",
    "配送体积",
    "装载率",
]


# SQLite INTEGER is signed 64-bit; values outside this range raise OverflowError on flush.
_SQLITE_INT_MIN = -(2**63)
_SQLITE_INT_MAX = 2**63 - 1


def _parse_date(value):
    if hasattr(value, "date"):
        return value.date()
    return datetime.strptime(str(value), "%Y-%m-%d").date()


def _parse_est_duration(value) -> int:
    duration = int(value or 0)
    if duration < _SQLITE_INT_MIN or duration > _SQLITE_INT_MAX:
        raise ValueError("预计时效超出有效整数范围")
    return duration


def import_excel(db: Session, dataset_type: str, file_path: str):
    wb = load_workbook(file_path)
    sheet = wb.active
    headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]
    header_map = {name: idx for idx, name in enumerate(headers)}

    errors = []
    total_rows = 0
    success_rows = 0

    for col in REQUIRED_COLUMNS:
        if col not in header_map:
            errors.append({"row": 1, "reason": f"缺少必填列: {col}"})
    if errors:
        return {"total_rows": 0, "success_rows": 0, "failed_rows": 0, "errors": errors}

    for row_no in range(2, sheet.max_row + 1):
        total_rows += 1
        values = [sheet.cell(row=row_no, column=i + 1).value for i in range(len(headers))]
        try:
            route_date = _parse_date(values[header_map["排线日期"]])
            waybill_no = str(values[header_map["运单号"]]).strip()
            route_line = str(values[header_map["归属线路"]]).strip()
            warehouse_name = str(values[header_map["始发仓库"]]).strip()
            stores = str(values[header_map["拼载门店"]]).strip()
            volume = float(values[header_map["配送体积"]])
            load_rate = float(str(values[header_map["装载率"]]).replace("%", ""))

            if not all([waybill_no, route_line, warehouse_name, stores]):
                raise ValueError("文本字段存在空值")
            if volume < 0:
                raise ValueError("配送体积不能为负数")

            if dataset_type == "system":
                est_distance = float(values[header_map.get("预计公里数", -1)] or 0)
                est_duration = _parse_est_duration(values[header_map.get("预计时效", -1)])
                obj = SysSuggest(
                    route_date=route_date,
                    waybill_no=waybill_no,
                    route_line=route_line,
                    warehouse_name=warehouse_name,
                    stores=stores,
                    volume=volume,
                    load_rate=load_rate,
                    est_distance=est_distance,
                    est_duration=est_duration,
                )
            else:
                obj = ManualRoute(
                    route_date=route_date,
                    waybill_no=waybill_no,
                    route_line=route_line,
                    warehouse_name=warehouse_name,
                    stores=stores,
                    volume=volume,
                    load_rate=load_rate,
                )
            db.add(obj)
            success_rows += 1
        except Exception as exc:
            errors.append({"row": row_no, "reason": str(exc)})

    db.commit()
    return {
        "total_rows": total_rows,
        "success_rows": success_rows,
        "failed_rows": total_rows - success_rows,
        "errors": errors,
    }
