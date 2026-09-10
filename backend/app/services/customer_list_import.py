"""从「客户列表」类 Excel 导入客户基础资料，并用收货坐标更新门店经纬度表。"""

from __future__ import annotations

import math
import re
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from sqlalchemy.orm import Session

from app.models.entities import CustomerProfile
from app.services.store_master_service import upsert_store_coordinate
from app.utils.finite import parse_finite_or_none

_REQUIRED = ("客户代码", "客户名称", "收货坐标")


def _cell_str(ws, r: int, c: int) -> str:
    v = ws.cell(r, c).value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if isinstance(v, float) and v == int(v) and (v == v) and (v not in (float("inf"), float("-inf"))):
            return str(int(v))
        return str(v)
    return str(v).strip()


def _cell_float_opt(ws, r: int, c: int) -> Optional[float]:
    return parse_finite_or_none(ws.cell(r, c).value)


def parse_lnglat(s: str) -> Optional[Tuple[float, float]]:
    t = (s or "").strip()
    if not t:
        return None
    t = t.replace("，", ",").replace("；", ";")
    m = re.match(
        r"^\s*(-?\d+(?:\.\d+)?)\s*[,;]\s*(-?\d+(?:\.\d+)?)\s*$", t, re.S
    )
    if m:
        try:
            lng, lat = float(m.group(1)), float(m.group(2))
        except (ValueError, TypeError):
            return None
    else:
        parts = t.replace(";", " ").split()
        if len(parts) < 2:
            parts = [p for p in t.split(",") if p.strip()]
        if len(parts) < 2:
            return None
        try:
            lng, lat = float(parts[0].strip()), float(parts[1].strip())
        except (ValueError, TypeError):
            return None
    if not (math.isfinite(lng) and math.isfinite(lat)):
        return None
    if not (-180 <= lng <= 180 and -90 <= lat <= 90):
        return None
    return (lng, lat)


def import_customer_list_workbook(
    db: Session, path: str, *, import_source: str, sheet_name: Optional[str] = None
) -> Dict[str, Any]:
    """
    导入客户基础资料到 customer_profile，并对有「收货坐标」且可解析的「客户名称」
    执行 store_coordinate upsert（data_source 为 import_source）。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        elif "客户列表" in wb.sheetnames:
            ws = wb["客户列表"]
        else:
            ws = wb[wb.sheetnames[0]]
        max_col = int(ws.max_column or 0)
        if max_col < 2:
            raise ValueError("工作表无有效表头或内容")

        col: Dict[str, int] = {}
        for c in range(1, max_col + 1):
            h = ws.cell(1, c).value
            if h is not None and str(h).strip():
                col[str(h).strip()] = c
        for need in _REQUIRED:
            if need not in col:
                head_preview = list(col.keys())[:25]
                raise ValueError(
                    f"缺少必需列「{need}」。已识别表头: {', '.join(head_preview)}"
                )

        tag = (import_source or "客户列表导入").strip() or "客户列表导入"
        n_rows = 0
        n_cp = 0
        n_coord_new = 0
        n_coord_upd = 0
        n_coord_skip = 0
        n_coord_err = 0
        errors: List[Dict[str, Any]] = []

        for r in range(2, (ws.max_row or 0) + 1):
            code = _cell_str(ws, r, col["客户代码"])
            name0 = _cell_str(ws, r, col["客户名称"])
            if not code and not name0:
                continue
            n_rows += 1
            if not code:
                n_coord_err += 1
                if len(errors) < 30:
                    errors.append({"row": r, "reason": "客户代码为空"})
                continue

            def cell_field(field: str) -> str:
                c1 = col.get(field)
                return _cell_str(ws, r, c1) if c1 else ""

            def text_f(field: str) -> Optional[str]:
                c1 = col.get(field)
                if not c1:
                    return None
                s = _cell_str(ws, r, c1)
                return s if s else None

            row = (
                db.query(CustomerProfile)
                .filter(CustomerProfile.customer_code == code)
                .first()
            )
            if not row:
                row = CustomerProfile(customer_code=code)
                db.add(row)
            # 字段写入
            row.customer_name = cell_field("客户名称") or code
            row.customer_type = text_f("客户类型")
            row.business_status = text_f("营业状态")
            row.contact_name = text_f("联系人")
            row.business_hours = text_f("营业时间")
            row.contact_phone = text_f("联系电话")
            row.province = text_f("省")
            row.city = text_f("市")
            row.district = text_f("区")
            row.county = text_f("县")
            if col.get("地址"):
                v = _cell_str(ws, r, col["地址"])
                row.address = v if v else None
            else:
                row.address = None
            row.performance_sla = text_f("履约时效")
            row.line_count = text_f("线路数")
            row.route_line = text_f("所属线路")
            row.e_sign = text_f("开启电子签")
            row.latest_delivery = text_f("最晚送达时间")
            row.auth_status = text_f("认证状态")
            if col.get("结算仓店距离（km）"):
                row.settlement_warehouse_km = _cell_float_opt(
                    ws, r, col["结算仓店距离（km）"]
                )
            if col.get("仓店距离（km）"):
                row.warehouse_store_km = _cell_float_opt(ws, r, col["仓店距离（km）"])
            row.customer_coordinate = text_f("客户坐标")
            row.driver_coordinate = text_f("司机上报坐标")
            row.carrier = text_f("所属承运商")
            if col.get("员工认证明细"):
                v = _cell_str(ws, r, col["员工认证明细"])
                row.staff_auth_detail = v if v else None
            row.customer_group = text_f("所属客户组")
            row.group_min_order = text_f("客户组起送量")
            row.min_order_type = text_f("起送量类型")
            row.min_order = text_f("起送量")
            row.delivery_schedule = text_f("配送排程")
            row.loop_mode = text_f("循环模式")
            if col.get("客户备注"):
                v = _cell_str(ws, r, col["客户备注"])
                row.remark = v if v else None
            row.consignee = text_f("收货人")
            row.delivery_coordinate = text_f("收货坐标")
            row.consignee_phone = text_f("收货电话")
            if col.get("收货地址"):
                v = _cell_str(ws, r, col["收货地址"])
                row.delivery_address = v if v else None
            row.delivery_province = text_f("收货省")
            row.delivery_city = text_f("收货市")
            row.delivery_district = text_f("收货区")
            row.delivery_street = text_f("收货街道")
            row.import_source = tag
            n_cp += 1

            store = (row.customer_name or "").strip()
            coord_raw = _cell_str(ws, r, col["收货坐标"])
            if not store:
                n_coord_skip += 1
                continue
            if not coord_raw:
                n_coord_skip += 1
                continue
            pt = parse_lnglat(coord_raw)
            if not pt:
                n_coord_err += 1
                if len(errors) < 30:
                    errors.append(
                        {"row": r, "reason": f"收货坐标无法解析: {coord_raw[:80]}"}
                    )
                continue
            lng, lat = pt
            try:
                _, op = upsert_store_coordinate(
                    db,
                    store_name=store,
                    longitude=lng,
                    latitude=lat,
                    data_source=tag,
                    commit=False,
                )
                if op == "inserted":
                    n_coord_new += 1
                else:
                    n_coord_upd += 1
            except ValueError as e:
                n_coord_err += 1
                if len(errors) < 30:
                    errors.append({"row": r, "reason": str(e)})

        db.commit()
    finally:
        wb.close()

    return {
        "message": "import finished",
        "data_rows": n_rows,
        "customer_profile_upserted": n_cp,
        "store_coordinate_inserted": n_coord_new,
        "store_coordinate_updated": n_coord_upd,
        "store_coordinate_skipped": n_coord_skip,
        "store_coordinate_error_rows": n_coord_err,
        "errors_sample": errors,
    }
