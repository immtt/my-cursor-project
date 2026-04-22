"""客户基础资料 `customer_profile` 的增删改查与 Excel 导出（列与 `customer_list_import` 一致）。"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.models.entities import CustomerProfile

# 与 `customer_list_import` 中写入字段对应的表头顺序一致，便于用导出文件再导入
_CUSTOMER_EXCEL_HEADER_ATTRS: List[Tuple[str, str]] = [
    ("客户代码", "customer_code"),
    ("客户名称", "customer_name"),
    ("客户类型", "customer_type"),
    ("营业状态", "business_status"),
    ("联系人", "contact_name"),
    ("营业时间", "business_hours"),
    ("联系电话", "contact_phone"),
    ("省", "province"),
    ("市", "city"),
    ("区", "district"),
    ("县", "county"),
    ("地址", "address"),
    ("履约时效", "performance_sla"),
    ("线路数", "line_count"),
    ("所属线路", "route_line"),
    ("开启电子签", "e_sign"),
    ("最晚送达时间", "latest_delivery"),
    ("认证状态", "auth_status"),
    ("结算仓店距离（km）", "settlement_warehouse_km"),
    ("仓店距离（km）", "warehouse_store_km"),
    ("客户坐标", "customer_coordinate"),
    ("司机上报坐标", "driver_coordinate"),
    ("所属承运商", "carrier"),
    ("员工认证明细", "staff_auth_detail"),
    ("所属客户组", "customer_group"),
    ("客户组起送量", "group_min_order"),
    ("起送量类型", "min_order_type"),
    ("起送量", "min_order"),
    ("配送排程", "delivery_schedule"),
    ("循环模式", "loop_mode"),
    ("客户备注", "remark"),
    ("收货人", "consignee"),
    ("收货坐标", "delivery_coordinate"),
    ("收货电话", "consignee_phone"),
    ("收货地址", "delivery_address"),
    ("收货省", "delivery_province"),
    ("收货市", "delivery_city"),
    ("收货区", "delivery_district"),
    ("收货街道", "delivery_street"),
]


def _dt_iso(v: Optional[datetime]) -> Optional[str]:
    if v is None:
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def customer_profile_to_item(row: CustomerProfile) -> Dict[str, Any]:
    return {
        "id": row.id,
        "customer_code": row.customer_code,
        "customer_name": row.customer_name,
        "customer_short_name": row.customer_short_name,
        "customer_category": row.customer_category,
        "sales_org": row.sales_org,
        "addr_street": row.addr_street,
        "settlement_unit": row.settlement_unit,
        "status": row.status,
        "created_by": row.created_by,
        "updated_by": row.updated_by,
        "customer_type": row.customer_type,
        "business_status": row.business_status,
        "contact_name": row.contact_name,
        "business_hours": row.business_hours,
        "contact_phone": row.contact_phone,
        "province": row.province,
        "city": row.city,
        "district": row.district,
        "county": row.county,
        "address": row.address,
        "performance_sla": row.performance_sla,
        "line_count": row.line_count,
        "route_line": row.route_line,
        "e_sign": row.e_sign,
        "latest_delivery": row.latest_delivery,
        "auth_status": row.auth_status,
        "settlement_warehouse_km": row.settlement_warehouse_km,
        "warehouse_store_km": row.warehouse_store_km,
        "customer_coordinate": row.customer_coordinate,
        "driver_coordinate": row.driver_coordinate,
        "carrier": row.carrier,
        "staff_auth_detail": row.staff_auth_detail,
        "customer_group": row.customer_group,
        "group_min_order": row.group_min_order,
        "min_order_type": row.min_order_type,
        "min_order": row.min_order,
        "delivery_schedule": row.delivery_schedule,
        "loop_mode": row.loop_mode,
        "remark": row.remark,
        "consignee": row.consignee,
        "delivery_coordinate": row.delivery_coordinate,
        "consignee_phone": row.consignee_phone,
        "delivery_address": row.delivery_address,
        "delivery_province": row.delivery_province,
        "delivery_city": row.delivery_city,
        "delivery_district": row.delivery_district,
        "delivery_street": row.delivery_street,
        "import_source": row.import_source,
        "created_at": _dt_iso(row.created_at),
        "updated_at": _dt_iso(row.updated_at),
    }


def _apply_payload(row: CustomerProfile, data: Dict[str, Any], *, for_create: bool) -> None:
    str_fields = (
        "customer_short_name",
        "customer_category",
        "sales_org",
        "addr_street",
        "settlement_unit",
        "status",
        "created_by",
        "updated_by",
        "customer_type",
        "business_status",
        "contact_name",
        "business_hours",
        "contact_phone",
        "province",
        "city",
        "district",
        "county",
        "address",
        "performance_sla",
        "line_count",
        "route_line",
        "e_sign",
        "latest_delivery",
        "auth_status",
        "customer_coordinate",
        "driver_coordinate",
        "carrier",
        "staff_auth_detail",
        "customer_group",
        "group_min_order",
        "min_order_type",
        "min_order",
        "delivery_schedule",
        "loop_mode",
        "remark",
        "consignee",
        "delivery_coordinate",
        "consignee_phone",
        "delivery_address",
        "delivery_province",
        "delivery_city",
        "delivery_district",
        "delivery_street",
        "import_source",
    )
    if for_create or "customer_code" in data:
        code = (data.get("customer_code") or "").strip()
        row.customer_code = code
    if for_create or "customer_name" in data:
        name = (data.get("customer_name") or "").strip()
        row.customer_name = name
    for k in str_fields:
        if k in data:
            v = data[k]
            if v is None or (isinstance(v, str) and not v.strip()):
                setattr(row, k, None)
            else:
                setattr(row, k, str(v).strip() if not isinstance(v, (int, float)) else str(v))
    for k in ("settlement_warehouse_km", "warehouse_store_km"):
        if k in data:
            v = data[k]
            if v is None or v == "":
                setattr(row, k, None)
            else:
                setattr(row, k, float(v))


def list_customer_profiles(
    db: Session,
    *,
    search: str = "",
    skip: int = 0,
    limit: int = 20,
) -> Tuple[int, List[CustomerProfile]]:
    if skip < 0:
        skip = 0
    limit = min(max(1, limit), 200)
    q = db.query(CustomerProfile)
    s = (search or "").strip()
    if s:
        like = f"%{s}%"
        q = q.filter(
            or_(
                CustomerProfile.customer_code.like(like),
                CustomerProfile.customer_name.like(like),
                CustomerProfile.customer_short_name.like(like),
                CustomerProfile.sales_org.like(like),
            )
        )
    total = q.count()
    items = (
        q.order_by(CustomerProfile.id.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return total, items


def get_customer_profile(db: Session, row_id: int) -> Optional[CustomerProfile]:
    return (
        db.query(CustomerProfile)
        .filter(CustomerProfile.id == row_id)
        .first()
    )


def get_by_customer_code(db: Session, customer_code: str) -> Optional[CustomerProfile]:
    return (
        db.query(CustomerProfile)
        .filter(CustomerProfile.customer_code == customer_code)
        .first()
    )


def create_customer_profile(db: Session, data: Dict[str, Any]) -> CustomerProfile:
    code = (data.get("customer_code") or "").strip()
    if not code:
        raise ValueError("客户编码不能为空")
    name = (data.get("customer_name") or "").strip()
    if not name:
        name = code
    if get_by_customer_code(db, code):
        raise ValueError("客户编码已存在")
    row = CustomerProfile(customer_code=code, customer_name=name)
    db.add(row)
    _apply_payload(row, data, for_create=True)
    row.customer_name = (row.customer_name or "").strip() or code
    db.commit()
    db.refresh(row)
    return row


def update_customer_profile(db: Session, row_id: int, data: Dict[str, Any]) -> Optional[CustomerProfile]:
    row = get_customer_profile(db, row_id)
    if not row:
        return None
    if "customer_code" in data:
        new_code = (data.get("customer_code") or "").strip()
        if not new_code:
            raise ValueError("客户编码不能为空")
        other = get_by_customer_code(db, new_code)
        if other and other.id != row_id:
            raise ValueError("客户编码已存在")
    _apply_payload(row, data, for_create=False)
    if (row.customer_name or "").strip() == "":
        row.customer_name = (row.customer_code or "").strip() or "—"
    db.commit()
    db.refresh(row)
    return row


def delete_customer_profile(db: Session, row_id: int) -> bool:
    row = get_customer_profile(db, row_id)
    if not row:
        return False
    db.delete(row)
    db.commit()
    return True


def build_customer_profile_export_xlsx_bytes(db: Session) -> bytes:
    """导出全部客户行；表头与导入模板一致，工作表名「客户列表」。"""
    rows = (
        db.query(CustomerProfile)
        .order_by(CustomerProfile.id.asc())
        .all()
    )
    wb: Workbook = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "客户列表"
    headers = [h for h, _ in _CUSTOMER_EXCEL_HEADER_ATTRS]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, value=h)
    for r, row in enumerate(rows, 2):
        for c, (_, attr) in enumerate(_CUSTOMER_EXCEL_HEADER_ATTRS, 1):
            v = getattr(row, attr, None)
            if v is None:
                ws.cell(r, c, value="")
            elif isinstance(v, float):
                ws.cell(r, c, value=v)
            else:
                ws.cell(r, c, value=str(v))
    buf = BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()
