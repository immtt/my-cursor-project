import os
import tempfile

import openpyxl

from app.models.entities import CustomerProfile, StoreCoordinate
from app.services.customer_list_import import import_customer_list_workbook, parse_lnglat


def test_parse_lnglat():
    assert parse_lnglat("116.1,31.2") == (116.1, 31.2)
    assert parse_lnglat("116.1，31.2") == (116.1, 31.2)
    assert parse_lnglat("") is None
    assert parse_lnglat("x") is None


def _build_min_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客户列表"
    headers = [
        "客户代码",
        "客户名称",
        "客户类型",
        "营业状态",
        "联系人",
        "营业时间",
        "联系电话",
        "省",
        "市",
        "区",
        "县",
        "地址",
        "履约时效",
        "线路数",
        "所属线路",
        "开启电子签",
        "最晚送达时间",
        "认证状态",
        "结算仓店距离（km）",
        "仓店距离（km）",
        "客户坐标",
        "司机上报坐标",
        "所属承运商",
        "员工认证明细",
        "所属客户组",
        "客户组起送量",
        "起送量类型",
        "起送量",
        "配送排程",
        "循环模式",
        "客户备注",
        "收货人",
        "收货坐标",
        "收货电话",
        "收货地址",
        "收货省",
        "收货市",
        "收货区",
        "收货街道",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    hmap = {h: i for i, h in enumerate(headers, start=1)}
    r2 = [""] * len(headers)
    r2[0] = "C1"
    r2[1] = "测试门店A"
    r2[hmap["地址"] - 1] = "某地址A"
    r2[hmap["收货坐标"] - 1] = "116.5,31.5"
    ws.append(r2)
    r3 = [""] * len(headers)
    r3[0] = "C2"
    r3[1] = "测试门店B"
    ws.append(r3)
    wb.save(path)
    wb.close()


def test_import_upsert_customer_and_store(db_session):
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        _build_min_workbook(path)
        r = import_customer_list_workbook(
            db_session, path, import_source="单元测试", sheet_name="客户列表"
        )
        assert r["customer_profile_upserted"] == 2
        assert r["store_coordinate_inserted"] == 1
        assert r["store_coordinate_skipped"] >= 1
        a = (
            db_session.query(CustomerProfile)
            .filter(CustomerProfile.customer_code == "C1")
            .one()
        )
        assert a.customer_name == "测试门店A"
        sc = (
            db_session.query(StoreCoordinate)
            .filter(StoreCoordinate.store_name == "测试门店A")
            .first()
        )
        assert sc is not None
        assert abs(sc.longitude - 116.5) < 0.01
        assert abs(sc.latitude - 31.5) < 0.01
        assert "单元测试" in (sc.data_source or "")
    finally:
        if os.path.exists(path):
            os.remove(path)


def _workbook_with_rows(path: str, data_rows: list[tuple]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客户列表"
    ws.append(["客户代码", "客户名称", "收货坐标"])
    for code, name, coord in data_rows:
        ws.append([code, name, coord])
    wb.save(path)
    wb.close()


def test_import_duplicate_customer_code_without_coords_last_wins(db_session):
    """Two rows for the same code, neither with a coordinate, must not abort the file.

    The first row is still pending (no flush) when the second is seen. Inserting
    a second CustomerProfile hits the unique constraint and rolls back every
    other valid customer in the workbook.
    """
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        _workbook_with_rows(
            path,
            [
                ("C1", "旧名称", None),
                ("C1", "新名称", None),
                ("C2", "其他客户", "121.47,31.23"),
            ],
        )
        r = import_customer_list_workbook(
            db_session, path, import_source="单元测试", sheet_name="客户列表"
        )
        assert r["customer_profile_upserted"] == 3
        rows = (
            db_session.query(CustomerProfile)
            .filter(CustomerProfile.customer_code.in_(["C1", "C2"]))
            .order_by(CustomerProfile.customer_code)
            .all()
        )
        assert [row.customer_code for row in rows] == ["C1", "C2"]
        assert rows[0].customer_name == "新名称"
        assert rows[1].customer_name == "其他客户"
        sc = (
            db_session.query(StoreCoordinate)
            .filter(StoreCoordinate.store_name == "其他客户")
            .one()
        )
        assert abs(sc.longitude - 121.47) < 0.01
    finally:
        if os.path.exists(path):
            os.remove(path)


def test_import_duplicate_customer_code_later_row_keeps_coordinates(db_session):
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        _workbook_with_rows(
            path,
            [
                ("C1", "旧名称", None),
                ("C1", "新名称", "116.4,39.9"),
            ],
        )
        r = import_customer_list_workbook(
            db_session, path, import_source="单元测试", sheet_name="客户列表"
        )
        assert r["customer_profile_upserted"] == 2
        assert r["store_coordinate_inserted"] == 1
        row = (
            db_session.query(CustomerProfile)
            .filter(CustomerProfile.customer_code == "C1")
            .one()
        )
        assert row.customer_name == "新名称"
        assert row.delivery_coordinate == "116.4,39.9"
        sc = (
            db_session.query(StoreCoordinate)
            .filter(StoreCoordinate.store_name == "新名称")
            .one()
        )
        assert abs(sc.longitude - 116.4) < 0.01
        assert (
            db_session.query(StoreCoordinate)
            .filter(StoreCoordinate.store_name == "旧名称")
            .first()
            is None
        )
    finally:
        if os.path.exists(path):
            os.remove(path)
