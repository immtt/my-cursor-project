import json
import os
import tempfile
from datetime import date

import pytest
from openpyxl import Workbook

from app.models.entities import ManualRoute, SysSuggest
from app.services.gaode_service import backfill_manual_routes_by_batch
from app.services.import_service import (
    fetch_active_import_page,
    fetch_import_batch_page,
    fetch_import_batch_rows,
    import_excel,
)


def _create_system_excel(path: str, volume: float):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    ws.append(["2026-04-20", "SYS001", "线路A", "仓库1", "门店甲,门店乙", "4.2米标箱", volume, "70%", 30, 60])
    wb.save(path)


def _create_system_excel_many(path: str, n: int):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    for i in range(n):
        ws.append(
            [
                "2026-04-20",
                f"SYS{i:03d}",
                "线路A",
                "仓库1",
                "门店甲",
                "4.2米标箱",
                float(i + 1),
                "70%",
                30,
                60,
            ]
        )
    wb.save(path)


def _create_system_excel_three_dates(path: str):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    rows = [
        ["2026-04-18", "D1", "线路A", "仓库1", "门店甲", "4.2米标箱", 1.0, "70%", 30, 60],
        ["2026-04-20", "D2", "线路A", "仓库1", "门店甲", "4.2米标箱", 2.0, "70%", 30, 60],
        ["2026-04-25", "D3", "线路A", "仓库1", "门店甲", "4.2米标箱", 3.0, "70%", 30, 60],
    ]
    for r in rows:
        ws.append(r)
    wb.save(path)


def _seed_sys_suggest(
    db_session, waybill: str, route_d: date, volume: float = 1.0, batch_id: str = "pre"
) -> None:
    db_session.add(
        SysSuggest(
            route_date=route_d,
            waybill_no=waybill,
            route_line="线路A",
            warehouse_name="仓库1",
            stores="某门店",
            vehicle_type="4.2米标箱",
            volume=volume,
            load_rate=70.0,
            est_distance=10.0,
            est_duration=30,
            batch_id=batch_id,
            is_active=1,
        )
    )
    db_session.commit()


def _seed_manual_route(
    db_session, waybill: str, route_d: date, volume: float = 1.0, batch_id: str = "pre"
) -> None:
    db_session.add(
        ManualRoute(
            route_date=route_d,
            waybill_no=waybill,
            route_line="线路A",
            warehouse_name="仓库1",
            stores="某门店",
            vehicle_type="4.2米标箱",
            volume=volume,
            load_rate=70.0,
            batch_id=batch_id,
            is_active=1,
            calc_status=0,
        )
    )
    db_session.commit()


def test_import_marks_same_day_other_system_rows_inactive(db_session):
    d = date(2026, 4, 20)
    _seed_sys_suggest(db_session, "OLD_WB", d)
    path = os.path.join(tempfile.gettempdir(), "sys_replace_day.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    ws.append(
        [d, "NEW_WB", "线路A", "仓库1", "门店甲", "4.2米标箱", 2.0, "70%", 30, 60],
    )
    wb.save(path)
    out = import_excel(db_session, "system", path, operator="t")
    assert out["same_day_deactivated"] == 1
    old = (
        db_session.query(SysSuggest)
        .filter(SysSuggest.route_date == d, SysSuggest.waybill_no == "OLD_WB")
        .one()
    )
    assert old.is_active == 0
    new = (
        db_session.query(SysSuggest)
        .filter(SysSuggest.route_date == d, SysSuggest.waybill_no == "NEW_WB")
        .one()
    )
    assert new.is_active == 1


def test_import_marks_same_day_other_manual_rows_inactive(db_session):
    d = date(2026, 4, 22)
    _seed_manual_route(db_session, "OLDM", d)
    path = os.path.join(tempfile.gettempdir(), "man_replace_day.xlsx")
    _create_manual_excel(path)
    # 覆盖 2026-04-22 的运单为 M001（与 _create_manual_excel 一致），OLDM 应被置 0
    out = import_excel(db_session, "manual", path, operator="t")
    assert out["success_rows"] == 1
    assert out["same_day_deactivated"] == 1
    old = (
        db_session.query(ManualRoute)
        .filter(ManualRoute.route_date == d, ManualRoute.waybill_no == "OLDM")
        .one()
    )
    assert old.is_active == 0
    cur = (
        db_session.query(ManualRoute)
        .filter(ManualRoute.route_date == d, ManualRoute.waybill_no == "M001")
        .one()
    )
    assert cur.is_active == 1


def test_manual_import_does_not_touch_system_rows_on_same_date(db_session):
    d = date(2026, 4, 22)
    _seed_sys_suggest(db_session, "SYSX", d)
    path = os.path.join(tempfile.gettempdir(), "man_only.xlsx")
    _create_manual_excel(path)
    out = import_excel(db_session, "manual", path, operator="t")
    assert out["same_day_deactivated"] == 0
    sys_row = (
        db_session.query(SysSuggest)
        .filter(SysSuggest.route_date == d, SysSuggest.waybill_no == "SYSX")
        .one()
    )
    assert sys_row.is_active == 1


def test_import_same_file_duplicate_waybill_last_wins_keeps_siblings(db_session):
    """同文件重复 (排线日期, 运单号) 不得 IntegrityError 整批回滚；末行覆盖，其它运单仍写入。"""
    path = os.path.join(tempfile.gettempdir(), "sys_dup_waybill.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    ws.append(["2026-04-20", "DUP001", "线路A", "仓库1", "门店甲", "4.2米标箱", 1.0, "70%", 10, 20])
    ws.append(["2026-04-20", "KEEP001", "线路A", "仓库1", "门店乙", "4.2米标箱", 2.0, "70%", 11, 21])
    ws.append(["2026-04-20", "DUP001", "线路B", "仓库2", "门店丙", "4.2米高栏", 9.0, "80%", 99, 88])
    wb.save(path)

    out = import_excel(db_session, "system", path, operator="t")
    assert out["failed_rows"] == 1
    assert out["success_rows"] == 2
    assert any("重复的排线日期+运单号" in (e.get("reason") or "") for e in out["errors"])

    rows = (
        db_session.query(SysSuggest)
        .filter(SysSuggest.route_date == date(2026, 4, 20), SysSuggest.is_active == 1)
        .all()
    )
    by_wb = {r.waybill_no: r for r in rows}
    assert set(by_wb) == {"DUP001", "KEEP001"}
    assert by_wb["DUP001"].volume == 9.0
    assert by_wb["DUP001"].route_line == "线路B"
    assert by_wb["DUP001"].warehouse_name == "仓库2"
    assert by_wb["DUP001"].est_distance == 99.0
    assert by_wb["DUP001"].est_duration == 88
    assert by_wb["KEEP001"].volume == 2.0


def test_import_manual_same_file_duplicate_waybill_last_wins(db_session):
    path = os.path.join(tempfile.gettempdir(), "man_dup_waybill.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
        ]
    )
    ws.append(["2026-04-22", "M001", "线路A", "仓库1", "门店甲", "4.2米标箱", 4.0, "70%"])
    ws.append(["2026-04-22", "M001", "线路C", "仓库3", "门店丁", "6.8米箱", 7.5, "65%"])
    wb.save(path)

    out = import_excel(db_session, "manual", path, operator="t")
    assert out["success_rows"] == 1
    assert out["failed_rows"] == 1
    rows = db_session.query(ManualRoute).filter(ManualRoute.route_date == date(2026, 4, 22)).all()
    assert len(rows) == 1
    assert rows[0].volume == 7.5
    assert rows[0].route_line == "线路C"
    assert rows[0].warehouse_name == "仓库3"


def test_import_upserts_same_waybill_and_route_date(db_session):
    file_a = os.path.join(tempfile.gettempdir(), "sys_a.xlsx")
    file_b = os.path.join(tempfile.gettempdir(), "sys_b.xlsx")
    _create_system_excel(file_a, 8.0)
    _create_system_excel(file_b, 9.0)

    import_excel(db_session, "system", file_a, operator="u1")
    import_excel(db_session, "system", file_b, operator="u1")

    rows = db_session.query(SysSuggest).filter(SysSuggest.route_date == date(2026, 4, 20)).all()
    active_rows = [r for r in rows if r.is_active == 1]
    assert len(rows) == 1
    assert len(active_rows) == 1
    assert active_rows[0].volume == 9.0


def test_fetch_import_batch_rows_matches_import_batch(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_rows.xlsx")
    _create_system_excel(path, 10.5)
    out = import_excel(db_session, "system", path, operator="t")
    bid = out["batch_id"]
    rows = fetch_import_batch_rows(db_session, "system", bid)
    assert len(rows) == 1
    assert rows[0]["waybill_no"] == "SYS001"
    assert rows[0]["volume"] == 10.5
    assert rows[0]["route_date"] == "2026-04-20"
    assert rows[0]["sort_order"] == 1


def test_fetch_import_batch_page_paginates(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_page.xlsx")
    _create_system_excel_many(path, 25)
    out = import_excel(db_session, "system", path, operator="t")
    bid = out["batch_id"]
    p1 = fetch_import_batch_page(db_session, "system", bid, 1, 20)
    assert p1["total"] == 25
    assert p1["page"] == 1
    assert p1["page_size"] == 20
    assert len(p1["items"]) == 20
    assert p1["items"][0]["waybill_no"] == "SYS000"
    assert p1["items"][0]["sort_order"] == 1
    assert p1["items"][19]["sort_order"] == 20
    p2 = fetch_import_batch_page(db_session, "system", bid, 2, 20)
    assert len(p2["items"]) == 5
    assert p2["items"][0]["waybill_no"] == "SYS020"
    assert p2["items"][0]["sort_order"] == 21
    assert p2["items"][-1]["sort_order"] == 25
    p50 = fetch_import_batch_page(db_session, "system", bid, 1, 50)
    assert len(p50["items"]) == 25


def test_fetch_import_batch_page_invalid_page_size(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_one.xlsx")
    _create_system_excel(path, 1.0)
    out = import_excel(db_session, "system", path, operator="t")
    bid = out["batch_id"]
    with pytest.raises(ValueError, match="page_size"):
        fetch_import_batch_page(db_session, "system", bid, 1, 10)


def test_api_import_batch_rejects_bad_page_size(client_api_authed):
    r = client_api_authed.get(
        "/api/import-batch",
        params={"dataset_type": "system", "batch_id": "x", "page_size": 10},
    )
    assert r.status_code == 400


def _create_system_excel_two_warehouses_same_date(path: str):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    ws.append(["2026-04-20", "W1", "线路A", "华东仓", "门店甲", "4.2米标箱", 1.0, "70%", 30, 60])
    ws.append(["2026-04-20", "W2", "线路A", "华北仓", "门店甲", "4.2米标箱", 2.0, "70%", 30, 60])
    wb.save(path)


def test_fetch_active_import_page_filters_by_route_date(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_dates.xlsx")
    _create_system_excel_three_dates(path)
    import_excel(db_session, "system", path, operator="t")
    mid = fetch_active_import_page(
        db_session, "system", date(2026, 4, 19), date(2026, 4, 21), 1, 20
    )
    assert mid["total"] == 1
    assert mid["items"][0]["waybill_no"] == "D2"
    single = fetch_active_import_page(
        db_session, "system", date(2026, 4, 18), date(2026, 4, 18), 1, 20
    )
    assert single["total"] == 1
    assert single["items"][0]["waybill_no"] == "D1"
    full = fetch_active_import_page(
        db_session, "system", date(2026, 4, 1), date(2026, 4, 30), 1, 50
    )
    assert full["total"] == 3


def test_fetch_active_import_page_filters_by_warehouse_and_lists_options(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_two_wh.xlsx")
    _create_system_excel_two_warehouses_same_date(path)
    import_excel(db_session, "system", path, operator="t")
    full = fetch_active_import_page(
        db_session, "system", date(2026, 4, 1), date(2026, 4, 30), 1, 20
    )
    assert set(full["warehouse_options"]) == {"华东仓", "华北仓"}
    east = fetch_active_import_page(
        db_session,
        "system",
        date(2026, 4, 1),
        date(2026, 4, 30),
        1,
        20,
        warehouse_name="华东仓",
    )
    assert east["total"] == 1
    assert east["items"][0]["waybill_no"] == "W1"


def test_fetch_active_import_page_filters_by_store_name_substring(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_two_stores_one_date.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    ws.append(
        ["2026-04-20", "S_A", "线路A", "仓1", "门店甲,门店乙", "4.2米标箱", 1.0, "70%", 30, 60]
    )
    ws.append(["2026-04-20", "S_B", "线路A", "仓1", "门店丙", "4.2米标箱", 1.0, "70%", 30, 60])
    wb.save(path)
    import_excel(db_session, "system", path, operator="t")
    all_rows = fetch_active_import_page(
        db_session, "system", date(2026, 4, 1), date(2026, 4, 30), 1, 20
    )
    assert all_rows["total"] == 2
    m = fetch_active_import_page(
        db_session,
        "system",
        date(2026, 4, 1),
        date(2026, 4, 30),
        1,
        20,
        None,
        "门店甲",
    )
    assert m["total"] == 1
    assert m["items"][0]["waybill_no"] == "S_A"
    c = fetch_active_import_page(
        db_session,
        "system",
        date(2026, 4, 1),
        date(2026, 4, 30),
        1,
        20,
        None,
        "门店丙",
    )
    assert c["total"] == 1
    assert c["items"][0]["waybill_no"] == "S_B"


def test_fetch_import_batch_page_filters_by_store_name(db_session):
    path = os.path.join(tempfile.gettempdir(), "sys_batch_store_filter.xlsx")
    _wb = Workbook()
    _ws = _wb.active
    _ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
            "预计公里数",
            "预计时效",
        ]
    )
    _ws.append(
        ["2026-04-20", "B_A", "线路A", "仓1", " Alpha店 ", "4.2米标箱", 1.0, "70%", 30, 60]
    )
    _ws.append(["2026-04-20", "B_B", "线路A", "仓1", "Beta店", "4.2米标箱", 1.0, "70%", 30, 60])
    _wb.save(path)
    out = import_excel(db_session, "system", path, operator="t")
    bid = out["batch_id"]
    p0 = fetch_import_batch_page(db_session, "system", bid, 1, 20, None)
    assert p0["total"] == 2
    p_alpha = fetch_import_batch_page(db_session, "system", bid, 1, 20, "Alpha")
    assert p_alpha["total"] == 1
    assert p_alpha["items"][0]["waybill_no"] == "B_A"


def test_fetch_active_import_page_rejects_inverted_range(db_session):
    with pytest.raises(ValueError, match="route_date_from"):
        fetch_active_import_page(
            db_session, "system", date(2026, 4, 25), date(2026, 4, 20), 1, 20
        )


def test_api_import_active_rejects_inverted_dates(client_api_authed):
    r = client_api_authed.get(
        "/api/import-active",
        params={
            "dataset_type": "system",
            "route_date_from": "2026-04-25",
            "route_date_to": "2026-04-20",
            "page": 1,
            "page_size": 20,
        },
    )
    assert r.status_code == 400


def _create_manual_excel(path: str):
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "排线日期",
            "运单号",
            "归属线路",
            "始发仓库",
            "拼载门店",
            "车辆类型",
            "配送体积",
            "装载率",
        ]
    )
    ws.append(["2026-04-22", "M001", "线路A", "仓库1", "门店甲", "4.2米标箱", 4.0, "70%"])
    wb.save(path)


def test_backfill_manual_routes_by_batch_sets_estimates(db_session):
    path = os.path.join(tempfile.gettempdir(), "manual_est.xlsx")
    _create_manual_excel(path)
    out = import_excel(db_session, "manual", path, operator="t")
    bid = out["batch_id"]
    row = db_session.query(ManualRoute).filter(ManualRoute.batch_id == bid).first()
    assert row is not None
    assert row.calc_status == 0
    res = backfill_manual_routes_by_batch(db_session, bid)
    db_session.refresh(row)
    assert res["updated"] == 1
    assert res["failed"] == 0
    assert row.est_distance is not None
    assert row.est_duration is not None
    assert row.calc_status == 1
    assert row.delivery_store_order
    assert json.loads(row.delivery_store_order) == ["门店甲"]


def test_api_manual_backfill_unknown_batch_returns_zero(client_api_authed):
    r = client_api_authed.post("/api/manual/backfill", json={"batch_id": "__no_such_batch__"})
    assert r.status_code == 200
    data = r.json()
    assert data["updated"] == 0
    assert data["failed"] == 0
    assert data["failures"] == []


def test_api_manual_backfill_rejects_neither_batch_nor_range(client_api_authed):
    r = client_api_authed.post("/api/manual/backfill", json={})
    assert r.status_code == 422


def test_api_manual_backfill_by_date_range_updates_zero_when_empty(client_api_authed):
    r = client_api_authed.post(
        "/api/manual/backfill",
        json={"route_date_from": "2099-01-01", "route_date_to": "2099-01-31"},
    )
    assert r.status_code == 200
    data = r.json()
    assert data["updated"] == 0
    assert data["failed"] == 0
