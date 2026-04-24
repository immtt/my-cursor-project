import os
import tempfile
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.main as app_main
from app.api.deps_auth import verify_bearer_token
from app.db.session import Base, get_db
from app.models import entities  # noqa: F401  — 注册表
from app.models.entities import AppUser
from app.services.auth_service import hash_password
from starlette.requests import Request


def _ephemeral_client():
    db_file = os.path.join(tempfile.gettempdir(), "smart_route_cp_api_test.db")
    if os.path.exists(db_file):
        os.remove(db_file)
    engine = create_engine(f"sqlite:///{db_file}", connect_args={"check_same_thread": False})
    TestingLocal = sessionmaker(bind=engine, autocommit=False, autoflush=False)
    Base.metadata.create_all(bind=engine)

    def _get_db():
        db = TestingLocal()
        try:
            yield db
        finally:
            db.close()

    def _verify_bypass(request: Request) -> None:
        if request.method == "OPTIONS":
            return
        request.state.user_id = 1

    app_main._core.dependency_overrides[get_db] = _get_db
    app_main._core.dependency_overrides[verify_bearer_token] = _verify_bypass
    seed = TestingLocal()
    try:
        seed.add(
            AppUser(
                username="cp_e2e",
                password_hash=hash_password("x"),
                is_active=True,
                is_admin=False,
            )
        )
        seed.commit()
    finally:
        seed.close()
    # 直接测 FastAPI 应用，避免 CORS 中间件对 response body 的边界问题
    client = TestClient(app_main._core)
    return client, db_file


def test_customer_profiles_crud_and_export():
    client, db_file = _ephemeral_client()
    try:
        r = client.get("/api/customer-profiles?limit=10")
        assert r.status_code == 200
        assert r.json()["total"] == 0

        r = client.post(
            "/api/customer-profiles",
            json={
                "customer_code": "K001",
                "customer_name": "测试客户",
                "city": "上海",
                "import_source": "单元测试",
            },
        )
        assert r.status_code == 200
        row = r.json()
        assert row["id"] >= 1
        rid = row["id"]
        assert row["customer_code"] == "K001"

        r2 = client.post(
            "/api/customer-profiles",
            json={"customer_code": "K001", "customer_name": "重复"},
        )
        assert r2.status_code == 400

        r3 = client.get(f"/api/customer-profiles/{rid}")
        assert r3.json()["customer_name"] == "测试客户"

        r4 = client.put(
            f"/api/customer-profiles/{rid}",
            json={"customer_name": "已改", "settlement_warehouse_km": 1.5},
        )
        assert r4.status_code == 200
        assert r4.json()["customer_name"] == "已改"
        assert r4.json()["settlement_warehouse_km"] == 1.5

        rx = client.get("/api/customer-profiles/export")
        assert rx.status_code == 200
        assert rx.content[:2] == b"PK", "应为 xlsx(zip) 流"
        wb = load_workbook(BytesIO(rx.content))
        assert "客户列表" in wb.sheetnames
        ws = wb["客户列表"]
        assert ws.max_row == 2
        assert ws.cell(1, 1).value == "客户代码"
        assert ws.cell(2, 1).value == "K001"
        assert ws.cell(2, 2).value == "已改"
        wb.close()

        d = client.delete(f"/api/customer-profiles/{rid}")
        assert d.status_code == 200
        d2 = client.get(f"/api/customer-profiles/{rid}")
        assert d2.status_code == 404
    finally:
        app_main._core.dependency_overrides.clear()
        if os.path.exists(db_file):
            os.remove(db_file)
