from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.import_service import build_import_template_xlsx


def test_build_import_template_system_columns():
    content, filename = build_import_template_xlsx("system")
    assert filename.endswith(".xlsx")
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert headers == [
        "排线日期",
        "运单号",
        "归属线路",
        "始发仓库",
        "拼载门店",
        "车辆类型",
        "配送体积",
        "装载率",
        "预计公里数",
        "预计时效",
    ]
    assert ws.max_row >= 2


def test_build_import_template_manual_columns():
    content, _ = build_import_template_xlsx("manual")
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "预计公里数" not in headers
    assert "预计时效" not in headers
    assert "配送体积" in headers


def test_build_import_template_invalid_type():
    with pytest.raises(ValueError):
        build_import_template_xlsx("other")


def test_api_import_template_get_ok():
    client = TestClient(app)
    r = client.get("/api/import-template", params={"dataset_type": "system"})
    assert r.status_code == 200
    assert "spreadsheetml" in (r.headers.get("content-type") or "")
    wb = load_workbook(BytesIO(r.content))
    assert wb.active.max_row >= 2
