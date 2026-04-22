#!/usr/bin/env python3
"""从「店间距离」类 Excel 导入 store_pair_distance 与 store_coordinate。

表头（每 sheet 第一行）：客户1 | 客户1坐标 | 客户2 | 客户2坐标 | 距离（km）
坐标格式：经度,纬度（如 117.152152,31.716447）

读取 backend/.env 的 DATABASE_URL；并执行 create_all 以创建新表。

用法：
  python3 scripts/import_store_distances.py /path/to/店间距离.xlsx
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
BACKEND = ROOT / "backend"
sys.path.insert(0, str(BACKEND))

from dotenv import load_dotenv
from openpyxl import load_workbook
from sqlalchemy.orm import Session

load_dotenv(BACKEND / ".env")

from app.db.session import SessionLocal, engine  # noqa: E402
from app.models.entities import Base, StoreCoordinate, StorePairDistance  # noqa: E402

# 确保新表存在
import app.models.entities  # noqa: F401, E402

Base.metadata.create_all(bind=engine)

HEADER_C1 = "客户1"
HEADER_C1_COORD = "客户1坐标"
HEADER_C2 = "客户2"
HEADER_C2_COORD = "客户2坐标"
HEADER_DIST = "距离（km）"


def _norm_cell(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return s


def parse_lng_lat(raw: str) -> tuple[float, float] | None:
    s = _norm_cell(raw)
    if not s:
        return None
    s = s.replace("，", ",")
    parts = [p.strip() for p in s.split(",") if p.strip()]
    if len(parts) < 2:
        m = re.match(r"^\s*([0-9.+-]+)\s+([0-9.+-]+)\s*$", s)
        if not m:
            return None
        parts = [m.group(1), m.group(2)]
    try:
        lng = float(parts[0])
        lat = float(parts[1])
    except ValueError:
        return None
    return lng, lat


def parse_distance_km(raw) -> float | None:
    s = _norm_cell(raw)
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def upsert_store_coord(
    db: Session,
    name: str,
    lng: float,
    lat: float,
    data_source: str,
) -> None:
    row = db.query(StoreCoordinate).filter(StoreCoordinate.store_name == name).first()
    if row:
        row.longitude = lng
        row.latitude = lat
        row.data_source = data_source
    else:
        db.add(
            StoreCoordinate(
                store_name=name,
                longitude=lng,
                latitude=lat,
                data_source=data_source,
            )
        )
        db.flush()


def import_workbook(path: Path, data_source_label: str) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    stats = {"sheets": 0, "distance_rows": 0, "stores_updated": 0, "skipped_rows": 0}
    db = SessionLocal()
    try:
        for sheet_name in wb.sheetnames:
            stats["sheets"] += 1
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if not header:
                continue
            header = [_norm_cell(h) for h in header]
            try:
                i1 = header.index(HEADER_C1)
                i1c = header.index(HEADER_C1_COORD)
                i2 = header.index(HEADER_C2)
                i2c = header.index(HEADER_C2_COORD)
                idist = header.index(HEADER_DIST)
            except ValueError:
                print(f"[skip] sheet {sheet_name!r}: 表头不匹配", file=sys.stderr)
                continue

            db.query(StorePairDistance).filter(StorePairDistance.region == sheet_name).delete(
                synchronize_session=False
            )

            region = sheet_name
            for row in rows_iter:
                if not row:
                    continue
                def cell(idx: int):
                    return row[idx] if idx < len(row) else None

                n1 = _norm_cell(cell(i1))
                n2 = _norm_cell(cell(i2))
                coord1 = parse_lng_lat(cell(i1c) or "")
                coord2 = parse_lng_lat(cell(i2c) or "")
                dist = parse_distance_km(cell(idist))

                if not n1 or not n2 or coord1 is None or coord2 is None or dist is None:
                    stats["skipped_rows"] += 1
                    continue

                upsert_store_coord(db, n1, coord1[0], coord1[1], data_source_label)
                upsert_store_coord(db, n2, coord2[0], coord2[1], data_source_label)

                db.add(
                    StorePairDistance(
                        region=region,
                        store_from=n1,
                        store_to=n2,
                        distance_km=dist,
                    )
                )
                stats["distance_rows"] += 1

            db.commit()

        # count distinct stores touched (approximate: query count)
        stats["stores_updated"] = db.query(StoreCoordinate).count()
        return stats
    finally:
        db.close()
        wb.close()


def main() -> None:
    ap = argparse.ArgumentParser(description="导入店间距离 Excel")
    ap.add_argument(
        "xlsx",
        nargs="?",
        default=str(Path.home() / "Downloads" / "店间距离.xlsx"),
        help="xlsx 路径（默认 ~/Downloads/店间距离.xlsx）",
    )
    ap.add_argument(
        "--source-label",
        default="店间距离.xlsx",
        help="写入 store_coordinate.data_source 的标记",
    )
    args = ap.parse_args()
    path = Path(os.path.expanduser(args.xlsx)).resolve()
    if not path.is_file():
        print(f"文件不存在: {path}", file=sys.stderr)
        sys.exit(1)
    out = import_workbook(path, args.source_label)
    print("导入完成:", out)
    print("DATABASE_URL:", engine.url.render_as_string(hide_password=True))


if __name__ == "__main__":
    main()
